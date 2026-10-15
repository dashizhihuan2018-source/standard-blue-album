import unittest
from pptx import Presentation
from pptx.util import Cm
from pptx.oxml.ns import qn
from references import b5_album_engine as engine


class BodyFontTest(unittest.TestCase):
    def render_run(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(5)).text_frame
        engine._render_block(frame, {'kind': 'para', 'text': '正文默认思源宋体。'}, True)
        return frame.paragraphs[0].runs[0]

    def assert_typefaces(self, run, expected):
        self.assertEqual(run.font.name, expected)
        props = run._r.get_or_add_rPr()
        for name in ('ea', 'cs'):
            self.assertEqual(props.find(qn('a:'+name)).get('typeface'), expected)

    def test_defaults_in_both_column_modes(self):
        for columns in (1, 2):
            with self.subTest(columns=columns), engine._options_scope({'columns': columns}):
                self.assert_typefaces(self.render_run(), 'Source Han Serif CN')

    def test_explicit_font_override_then_default_restored(self):
        with engine._options_scope({'body_font': 'Microsoft YaHei'}):
            self.assert_typefaces(self.render_run(), 'Microsoft YaHei')
        self.assert_typefaces(self.render_run(), 'Source Han Serif CN')

    def test_other_font_roles_unchanged(self):
        self.assertEqual(engine.SONG, 'Source Han Serif CN')
        self.assertEqual(engine.HEI, 'Source Han Serif CN')
        self.assertEqual(engine.TITLE_F, 'Source Han Serif CN')


if __name__ == '__main__':
    unittest.main()
