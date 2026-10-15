import copy
import unittest

from pptx import Presentation
from pptx.util import Cm
from references import b5_album_engine as engine


class HybridHeadingTest(unittest.TestCase):
    def blocks(self, title='9.1 从绘图进入 BIM 语义'):
        return [
            {'kind': 'chapter', 'text': '九、接口研究'},
            {'kind': 'heading', 'level': 2, 'text': title},
            {'kind': 'table', 'rows': [
                [{'text': '项目', 'images': []}, {'text': '概况', 'images': []}],
                [{'text': '平台', 'images': []}, {'text': '桌面软件', 'images': []}],
            ]},
            {'kind': 'para', 'text': '保留全部正文并从表格下方开始。' * 20},
        ]

    def test_heading_belongs_to_hybrid_not_a_header_gap(self):
        for columns in (1, 2):
            for size in ('B5', 'A4'):
                with self.subTest(columns=columns, size=size), engine._options_scope({'columns': columns, 'page_size': size}):
                    pages = engine.paginate(copy.deepcopy(self.blocks()))
                    self.assertEqual([p['type'] for p in pages], ['chapter', 'hybrid'])
                    page = pages[1]
                    self.assertEqual(page.get('lead_heading'), '9.1 从绘图进入 BIM 语义')
                    prs = Presentation()
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    engine.render_hybrid(slide, page, 2)
                    titles = [s for s in slide.shapes if s.has_text_frame and s.text == engine.head_disp(page['lead_heading'])]
                    self.assertEqual(len(titles), 1)
                    self.assertGreaterEqual(titles[0].top, Cm(engine.TOP))
                    table_label = next(s for s in slide.shapes if s.has_text_frame and s.text == '项目')
                    self.assertGreaterEqual(table_label.top, titles[0].top + titles[0].height)
                    below = next(s for s in slide.shapes if s.has_text_frame and '保留全部正文' in s.text)
                    self.assertGreater(below.top, table_label.top)
                    self.assertLessEqual(below.top + below.height, Cm(engine.CONTENT_BOTTOM) + 2)

    def test_long_heading_uses_more_space_without_losing_body(self):
        title = '9.1 ' + '含中英文 BIM 的多行标题测试' * 6
        pages = engine.paginate(copy.deepcopy(self.blocks(title)))
        hybrid = next(p for p in pages if p['type'] == 'hybrid')
        self.assertEqual(hybrid.get('lead_heading'), title)
        actual = ''.join(i['block']['text'] for p in pages for i in p.get('items', []) if i['block']['kind'] == 'para')
        self.assertEqual(actual, self.blocks()[-1]['text'])


if __name__ == '__main__':
    unittest.main()
